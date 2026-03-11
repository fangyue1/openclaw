import argparse
import json
from copy import copy
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--template', required=True)
    ap.add_argument('--cases', required=True)
    ap.add_argument('--out', required=True)
    args = ap.parse_args()

    cases = json.load(open(args.cases, 'r', encoding='utf-8'))
    wb = load_workbook(args.template)
    ws = wb[wb.sheetnames[0]]

    start_row = 8
    max_col = 11

    # clear old values, keep style and merged layout
    for r in range(start_row, ws.max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None

    style_src_row = start_row

    def clone_row_style(src, dst):
        ws.row_dimensions[dst].height = ws.row_dimensions[src].height
        for c in range(1, max_col + 1):
            s = ws.cell(src, c)
            d = ws.cell(dst, c)
            d._style = copy(s._style)
            d.number_format = s.number_format
            d.font = copy(s.font)
            d.fill = copy(s.fill)
            d.border = copy(s.border)
            d.alignment = copy(s.alignment)
            d.protection = copy(s.protection)

    def set_value(row, col, value):
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            for mr in ws.merged_cells.ranges:
                if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
                    ws.cell(mr.min_row, mr.min_col).value = value
                    return
            return
        cell.value = value

    for i, case in enumerate(cases):
        r = start_row + i
        if r > ws.max_row:
            ws.insert_rows(ws.max_row + 1, 1)
            clone_row_style(style_src_row, ws.max_row)

        set_value(r, 1, case.get('category') or 'コールバック予約')
        set_value(r, 2, '=ROW()-7')
        set_value(r, 3, case.get('title', ''))
        set_value(r, 4, case.get('steps', ''))
        set_value(r, 5, '')
        set_value(r, 6, case.get('precondition', ''))
        set_value(r, 7, case.get('expected', ''))
        set_value(r, 8, '')
        set_value(r, 9, '')
        set_value(r, 10, '')
        set_value(r, 11, '')

    wb.save(args.out)
    print(args.out)


if __name__ == '__main__':
    main()
