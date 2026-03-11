import argparse
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


HEADER_KEYWORDS = {
    "case_id": ["case", "caseid", "编号", "id", "No"],
    "category": ["category", "分類", "分类", "大分類"],
    "title": ["title", "标题", "試験項目", "测试项", "小分類"],
    "precondition": ["precondition", "前提", "前置", "テスト条件"],
    "steps": ["steps", "手順", "步骤", "確認手順"],
    "expected": ["expected", "期待", "预期", "想定結果"],
    "priority": ["priority", "优先", "優先"],
}

FALLBACK_ORDER = ["case_id", "category", "title", "precondition", "steps", "expected", "priority"]


def normalize(v):
    return str(v or "").strip().lower().replace(" ", "")


def detect_columns(ws):
    best_row, best_cols = 1, {}
    for row in ws.iter_rows(min_row=1, max_row=min(30, ws.max_row)):
        cols = {}
        for c in row:
            nv = normalize(c.value)
            for key, kws in HEADER_KEYWORDS.items():
                if key in cols:
                    continue
                if any(normalize(k) in nv for k in kws):
                    cols[key] = c.column
        if len(cols) > len(best_cols):
            best_cols = cols
            best_row = row[0].row
    return best_cols, best_row


def resolve_writable_pos(ws, row, col):
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return row, col
    # If cell is inside a merged range, write to top-left anchor of that range
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
            return mr.min_row, mr.min_col
    return row, col


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--cases", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    cases = json.loads(Path(args.cases).read_text(encoding="utf-8"))
    wb = load_workbook(args.template)
    ws = wb.active

    cols, header_row = detect_columns(ws)
    if not cols:
        cols = {k: i + 1 for i, k in enumerate(FALLBACK_ORDER)}
        header_row = 1
        for k, col in cols.items():
            ws.cell(row=1, column=col, value=k)

    start_row = header_row + 1
    for i, case in enumerate(cases, start=start_row):
        for key, col in cols.items():
            r, c = resolve_writable_pos(ws, i, col)
            ws.cell(row=r, column=c, value=case.get(key, ""))

    wb.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
