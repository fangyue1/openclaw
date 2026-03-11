import argparse
import json
from pathlib import Path
from openpyxl import load_workbook


def normalize(s):
    return str(s or "").strip()


def read_template_cases(template_path):
    wb = load_workbook(template_path, data_only=False)
    ws = wb[wb.sheetnames[0]]

    cases = []
    current_category = ""
    for r in range(8, ws.max_row + 1):
        c1 = normalize(ws.cell(r, 1).value)  # 大分類
        c3 = normalize(ws.cell(r, 3).value)  # 小分類
        c4 = normalize(ws.cell(r, 4).value)  # 確認手順
        c6 = normalize(ws.cell(r, 6).value)  # テスト条件
        c7 = normalize(ws.cell(r, 7).value)  # 想定結果

        if c1:
            current_category = c1

        if not any([c3, c4, c6, c7]):
            continue

        title = c3 or "（継続）"
        priority = "P1" if any(k in (title + c4 + c7) for k in ["異常", "エラー", "必須", "NG"]) else "P2"

        cases.append(
            {
                "case_id": f"TC-{len(cases)+1:04d}",
                "category": current_category or "コールバック予約",
                "title": title,
                "precondition": c6 or "ー",
                "steps": c4,
                "expected": c7,
                "priority": priority,
            }
        )
    return cases


def extract_ocr_lines(design_json):
    lines = []
    for o in design_json.get("ocr", []) or []:
        txt = normalize(o.get("text"))
        if not txt:
            continue
        for ln in txt.splitlines():
            s = normalize(ln)
            if len(s) >= 4:
                lines.append(s)

    uniq = []
    seen = set()
    for s in lines:
        k = s.lower()
        if k not in seen:
            seen.add(k)
            uniq.append(s)
    return uniq


def add_missing_from_ocr(cases, design_json, code_json):
    ocr_lines = extract_ocr_lines(design_json)
    base_text = "\n".join((c.get("title", "") + "\n" + c.get("steps", "") + "\n" + c.get("expected", "")) for c in cases)

    keywords = ["クリック", "入力", "確認", "遷移", "エラー", "必須", "表示", "ボタン", "メール", "予約"]
    added = 0
    for ln in ocr_lines:
        if added >= 20:
            break
        if not any(k in ln for k in keywords):
            continue
        if ln in base_text:
            continue

        cases.append(
            {
                "case_id": f"TC-{len(cases)+1:04d}",
                "category": "補足（設計書OCR）",
                "title": "画面文言/動作の補足確認",
                "precondition": "ー",
                "steps": f"設計書画像の記載に基づき、対象画面で次を確認する：{ln[:120]}",
                "expected": "設計書の記載どおりに表示・遷移・入力チェックが行えること",
                "priority": "P2",
            }
        )
        added += 1

    branches = int((code_json.get("summary") or {}).get("branches", 0) or 0)
    if branches > 0:
        cases.append(
            {
                "case_id": f"TC-{len(cases)+1:04d}",
                "category": "補足（コード観点）",
                "title": "分岐ロジック確認",
                "precondition": "条件を満たすデータを準備",
                "steps": "主要な分岐条件を満たす入力で操作する",
                "expected": "分岐結果が仕様どおりであること",
                "priority": "P1",
            }
        )

    return cases


def to_markdown(cases):
    lines = ["# Unit Test Cases (Template Skeleton Based)", "", "|No|Category|Title|Steps|Expected|", "|---|---|---|---|---|"]
    for i, c in enumerate(cases, 1):
        lines.append(f"|{i}|{c['category']}|{c['title']}|{c['steps'].replace('|','/')}|{c['expected'].replace('|','/')}|")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--design", required=True)
    ap.add_argument("--code", required=True)
    ap.add_argument("--out-cases", required=True)
    ap.add_argument("--out-md", required=True)
    args = ap.parse_args()

    design_json = json.loads(Path(args.design).read_text(encoding="utf-8"))
    code_json = json.loads(Path(args.code).read_text(encoding="utf-8"))

    cases = read_template_cases(args.template)
    cases = add_missing_from_ocr(cases, design_json, code_json)

    Path(args.out_cases).write_text(json.dumps(cases, ensure_ascii=False, indent=2), encoding="utf-8")
    Path(args.out_md).write_text(to_markdown(cases), encoding="utf-8")
    print(f"Wrote {args.out_cases} and {args.out_md}")


if __name__ == "__main__":
    main()
