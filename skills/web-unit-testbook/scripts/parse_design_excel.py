import argparse
import json
import re
from pathlib import Path
from zipfile import ZipFile
import os

from openpyxl import load_workbook

try:
    import pytesseract
    from PIL import Image
except Exception:
    pytesseract = None
    Image = None


def is_red(rgb) -> bool:
    if not rgb:
        return False
    # openpyxl color may be a string, RGB object, or theme/indexed color holder
    if not isinstance(rgb, str):
        rgb = getattr(rgb, "rgb", None) or getattr(rgb, "value", None)
    if not isinstance(rgb, str) or not rgb:
        return False
    rgb = rgb.upper().replace("#", "")
    if len(rgb) == 8:  # AARRGGBB -> RRGGBB
        rgb = rgb[2:]
    return rgb.startswith("FF") and rgb.endswith("0000")


def extract_check_rules(text: str):
    if not text:
        return []
    patterns = [
        ("required", r"必填|必須|required"),
        ("range", r"\d+\s*~\s*\d+"),
        ("maxmin", r"最大\d+|最小\d+|max\s*\d+|min\s*\d+"),
        ("email", r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+"),
    ]
    return [name for name, pattern in patterns if re.search(pattern, text, flags=re.IGNORECASE)]


def extract_embedded_images(xlsx_path: Path, media_dir: Path):
    media_dir.mkdir(parents=True, exist_ok=True)
    extracted = []
    with ZipFile(xlsx_path, "r") as zf:
        for name in zf.namelist():
            if name.startswith("xl/media/"):
                out = media_dir / Path(name).name
                out.write_bytes(zf.read(name))
                extracted.append(out)
    return extracted


def run_ocr(image_paths, lang):
    if not pytesseract or not Image:
        return []

    # Windows fallback: use default install path when PATH isn't refreshed yet
    win_tess = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if os.name == 'nt' and os.path.exists(win_tess):
        pytesseract.pytesseract.tesseract_cmd = win_tess

    results = []
    for i, p in enumerate(image_paths, start=1):
        try:
            txt = pytesseract.image_to_string(Image.open(p), lang=lang)
            results.append({"image_index": i, "image": str(p), "text": txt.strip()})
        except Exception as e:
            results.append({"image_index": i, "image": str(p), "error": str(e)})
    return results


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--ocr-lang", default="")
    args = ap.parse_args()

    xlsx = Path(args.input)
    wb = load_workbook(xlsx, data_only=True)

    fields = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            values = [c.value for c in row]
            joined = " | ".join([str(v) for v in values if v is not None])
            if not joined.strip():
                continue

            api_candidates = re.findall(r"[a-zA-Z_][a-zA-Z0-9_]*(?:__c|__r)?", joined)
            required = False
            cell_refs = []
            for c in row:
                if c.value is None:
                    continue
                cell_refs.append(c.coordinate)
                color = None
                try:
                    color = c.fill.start_color.rgb or c.font.color.rgb
                except Exception:
                    pass
                if color and is_red(color):
                    required = True

            fields.append(
                {
                    "sheet": ws.title,
                    "cells": cell_refs,
                    "raw": joined,
                    "api_names": sorted(set(api_candidates))[:10],
                    "required_hint": required,
                    "check_patterns": extract_check_rules(joined),
                }
            )

    media_dir = xlsx.parent / "_media"
    images = extract_embedded_images(xlsx, media_dir)
    ocr = run_ocr(images, args.ocr_lang) if args.ocr_lang else []

    result = {
        "source": str(xlsx),
        "sheet_count": len(wb.worksheets),
        "field_rules": fields,
        "embedded_images": [str(p) for p in images],
        "ocr": ocr,
    }

    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
