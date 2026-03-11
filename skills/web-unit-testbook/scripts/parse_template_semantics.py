import argparse
import json
from pathlib import Path
from openpyxl import load_workbook


def norm(v):
    return str(v or "").strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    wb = load_workbook(args.template)
    ws = wb[wb.sheetnames[0]]

    header_top = 6
    header_sub = 7

    cols = {}
    for c in range(1, ws.max_column + 1):
        v6 = norm(ws.cell(header_top, c).value)
        v7 = norm(ws.cell(header_sub, c).value)
        v = f"{v6}|{v7}"
        if v6 == "大分類":
            cols["category"] = c
        elif v6 in ("No.", "No"):
            cols["no"] = c
        elif v6 == "小分類":
            cols["title"] = c
        elif v6 == "確認手順":
            cols["steps"] = c
        elif v6 == "アカウント":
            cols["account"] = c
        elif v6 == "テスト条件":
            cols["condition"] = c
        elif v6 == "想定結果":
            cols["expected"] = c
        elif v6 == "テスト" and v7 == "結果":
            cols["test_result"] = c
        elif v6 == "テスト" and v7 == "実施者":
            cols["executor"] = c
        elif v6 == "テスト" and v7 == "実施日":
            cols["executed_at"] = c
        elif v6 == "テスト" and v7 == "備考":
            cols["note"] = c

    sample = {}
    for key, c in cols.items():
        sample[key] = norm(ws.cell(8, c).value)

    out = {
        "sheet": ws.title,
        "header_rows": [header_top, header_sub],
        "start_row": 8,
        "columns": cols,
        "sample_row_8": sample,
        "write_fields": ["category", "title", "steps", "condition", "expected"],
    }

    Path(args.out).write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
